#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
重新生成票务系统功能截图对照Word文档（清理404后）
"""

import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

MENU_DATA_FILE = "/workspace/menu_structure_v2.json"
SCREENSHOTS_DIR = "/workspace/screenshots_v2"
OUTPUT_DOC = "/workspace/票务系统功能截图对照文档.docx"

class WordDocGenerator:
    def __init__(self):
        self.doc = Document()
        self.screenshots_dir = Path(SCREENSHOTS_DIR)
        
        # 设置中文字体
        self.doc.styles['Normal'].font.name = 'Microsoft YaHei'
        self.doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        
        # 加载菜单数据
        with open(MENU_DATA_FILE, 'r', encoding='utf-8') as f:
            self.menu_data = json.load(f)
        
        print(f"加载了 {len(self.menu_data)} 条有效数据")
    
    def add_title(self, text, level=1):
        """添加标题"""
        if level == 0:
            heading = self.doc.add_heading(text, level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in heading.runs:
                run.font.name = 'Microsoft YaHei'
                run.font.size = Pt(22)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 51, 102)
        else:
            heading = self.doc.add_heading(text, level=level)
            for run in heading.runs:
                run.font.name = 'Microsoft YaHei'
                run.font.size = Pt(18) if level == 1 else Pt(14)
                run.font.color.rgb = RGBColor(0, 102, 204)
    
    def add_function_item(self, function_name, description, screenshot_file):
        """添加功能项：功能名称 + 描述 + 截图"""
        table = self.doc.add_table(rows=2, cols=1)
        table.style = 'Light Grid Accent 1'
        
        # 第一行：功能名称和描述
        cell = table.rows[0].cells[0]
        paragraph = cell.paragraphs[0]
        
        run = paragraph.add_run(function_name)
        run.font.name = 'Microsoft YaHei'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        
        if description and description != function_name:
            run = paragraph.add_run(f" - {description}")
            run.font.name = 'Microsoft YaHei'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(64, 64, 64)
        
        # 第二行：截图
        cell = table.rows[1].cells[0]
        screenshot_path = self.screenshots_dir / screenshot_file
        
        if screenshot_path.exists():
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                run = paragraph.add_run()
                run.add_picture(str(screenshot_path), width=Inches(6.0))
            except Exception as e:
                paragraph.add_run(f"[截图加载失败: {str(e)}]")
        else:
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run("[该功能截图暂时无法获取，可能需要特定权限]")
            run.font.name = 'Microsoft YaHei'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(150, 150, 150)
            run.font.italic = True
        
        self.doc.add_paragraph()
    
    def generate_document(self):
        """生成完整文档"""
        # 添加文档标题
        self.add_title("票务系统功能截图对照文档", level=0)
        
        # 添加文档说明
        p = self.doc.add_paragraph()
        p.add_run("本文档展示票务系统各功能模块的界面截图。注：部分功能可能需要特定权限才能访问。")
        run = p.runs[0]
        run.font.name = 'Microsoft YaHei'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(100, 100, 100)
        
        self.doc.add_paragraph()
        
        # 统计信息
        stats = {
            '票务管理': 0,
            '订单管理': 0,
            '工作台': 0,
            '通行管理': 0,
            '控制面板': 0,
            '报表管理': 0,
            '系统管理': 0,
        }
        
        # 统计各模块截图数量
        for item in self.menu_data:
            path = item['path']
            if '票务管理' in path:
                stats['票务管理'] += 1
            elif '订单管理' in path:
                stats['订单管理'] += 1
            elif '工作台' in path:
                stats['工作台'] += 1
            elif '通行管理' in path:
                stats['通行管理'] += 1
            elif '控制面板' in path:
                stats['控制面板'] += 1
            elif '报表管理' in path:
                stats['报表管理'] += 1
            elif '系统管理' in path:
                stats['系统管理'] += 1
        
        # 1. 票务管理
        self.add_title("一、票务管理", level=1)
        p = self.doc.add_paragraph("雪票电子纸质的全面管理")
        p.runs[0].font.name = 'Microsoft YaHei'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.italic = True
        
        for item in self.menu_data:
            if '票务管理' in item['path']:
                self.add_function_item(
                    item['path'].replace('票务管理-', ''),
                    item['description'],
                    item['screenshot']
                )
        
        # 2. 订单管理
        self.add_title("二、订单管理", level=1)
        p = self.doc.add_paragraph("对各类雪票、教务订单、租赁订单进行管理")
        p.runs[0].font.name = 'Microsoft YaHei'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.italic = True
        
        for item in self.menu_data:
            if '订单管理' in item['path']:
                self.add_function_item(
                    item['path'].replace('订单管理-', ''),
                    item['description'],
                    item['screenshot']
                )
        
        # 3. 工作台
        self.add_title("三、工作台", level=1)
        p = self.doc.add_paragraph("实现对现场售票及教练预约的下单管理")
        p.runs[0].font.name = 'Microsoft YaHei'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.italic = True
        
        for item in self.menu_data:
            if '工作台' in item['path']:
                self.add_function_item(
                    item['path'].replace('工作台-', ''),
                    item['description'],
                    item['screenshot']
                )
        
        # 4. 通行管理
        self.add_title("四、通行管理", level=1)
        p = self.doc.add_paragraph("通过闸机的管控实现对进出雪场的人员进行管理；通行方式可以是人脸或雪票二维码")
        p.runs[0].font.name = 'Microsoft YaHei'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.italic = True
        
        for item in self.menu_data:
            if '通行管理' in item['path']:
                self.add_function_item(
                    item['path'].replace('通行管理-', ''),
                    item['description'],
                    item['screenshot']
                )
        
        # 5. 控制面板
        self.add_title("五、控制面板", level=1)
        p = self.doc.add_paragraph("包括对雪票产品设置、电子/纸质票设置、闸机配置、售票小程序、教练、分销商、节假日、租赁物等运营参数的管理")
        p.runs[0].font.name = 'Microsoft YaHei'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.italic = True
        
        for item in self.menu_data:
            if '控制面板' in item['path']:
                self.add_function_item(
                    item['path'].replace('控制面板-', ''),
                    item['description'],
                    item['screenshot']
                )
        
        # 6. 报表管理
        self.add_title("六、报表管理", level=1)
        p = self.doc.add_paragraph("提供各类运营报表，包括票务核销报表、教学核销报表、教学数据看板等")
        p.runs[0].font.name = 'Microsoft YaHei'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.italic = True
        
        for item in self.menu_data:
            if '报表管理' in item['path']:
                self.add_function_item(
                    item['path'].replace('报表管理-', ''),
                    item['description'],
                    item['screenshot']
                )
        
        # 7. 系统管理
        self.add_title("七、系统管理", level=1)
        p = self.doc.add_paragraph("系统基础功能的设置，包括文件，身份标识，日志及相关系统参数的设置")
        p.runs[0].font.name = 'Microsoft YaHei'
        p.runs[0].font.size = Pt(11)
        p.runs[0].font.italic = True
        
        for item in self.menu_data:
            if '系统管理' in item['path']:
                self.add_function_item(
                    item['path'].replace('系统管理-', ''),
                    item['description'],
                    item['screenshot']
                )
        
        # 保存文档
        self.doc.save(OUTPUT_DOC)
        print(f"\n✓ Word文档已生成: {OUTPUT_DOC}")
        
        # 打印统计信息
        print("\n文档统计：")
        for module, count in stats.items():
            print(f"  {module}: {count} 个截图")

if __name__ == "__main__":
    print("=" * 60)
    print("重新生成Word对照文档...")
    print("=" * 60)
    
    generator = WordDocGenerator()
    generator.generate_document()
    
    print("\n" + "=" * 60)
    print("文档生成完成！")
    print(f"输出文件: {OUTPUT_DOC}")
    print("=" * 60)
