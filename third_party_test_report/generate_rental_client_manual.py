#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成租赁设备客户端发布手册
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUTPUT_FILE = "/workspace/third_party_test_report/租赁设备客户端发布手册.docx"

class RentalClientManualGenerator:
    def __init__(self):
        self.doc = Document()
        self.setup_styles()
    
    def setup_styles(self):
        """设置文档样式"""
        style = self.doc.styles['Normal']
        style.font.name = '微软雅黑'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    
    def add_title(self):
        """添加标题"""
        title = self.doc.add_heading('租赁设备客户端发布手册', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        title_run = title.runs[0]
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0, 51, 153)
        
        # 添加发布日期
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_run = p.add_run('发布日期：2023-11-05')
        date_run.font.size = Pt(14)
        date_run.font.color.rgb = RGBColor(68, 114, 196)
        
        self.doc.add_paragraph()
    
    def add_overview(self):
        """添加文档说明"""
        self.doc.add_heading('一、文档说明', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run('本手册用于指导租赁设备客户端应用程序的安装、配置和使用，确保租赁设备能够正常运行并与票务系统对接。')
        
        self.doc.add_paragraph()
        
        # 基本信息
        self.doc.add_heading('1.1 基本信息', level=2)
        
        table = self.doc.add_table(rows=10, cols=2)
        table.style = 'Light Grid Accent 1'
        
        table.rows[0].cells[0].text = '项目'
        table.rows[0].cells[1].text = '信息'
        
        for cell in table.rows[0].cells:
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        table.rows[1].cells[0].text = '应用名称'
        table.rows[1].cells[1].text = '租赁设备客户端'
        
        table.rows[2].cells[0].text = '适用平台'
        table.rows[2].cells[1].text = 'Android'
        
        table.rows[3].cells[0].text = '供应商'
        table.rows[3].cells[1].text = '亿思维科技有限公司'
        
        table.rows[4].cells[0].text = '设备品牌'
        table.rows[4].cells[1].text = 'Sunmi (商米)'
        
        table.rows[5].cells[0].text = '设备型号'
        table.rows[5].cells[1].text = 'V2S'
        
        table.rows[6].cells[0].text = '设备数量'
        table.rows[6].cells[1].text = '15台'
        
        table.rows[7].cells[0].text = '采购来源'
        table.rows[7].cells[1].text = '亿思维科技有限公司'
        
        table.rows[8].cells[0].text = '发布日期'
        table.rows[8].cells[1].text = '2023-11-05'
        
        table.rows[9].cells[0].text = '对接系统'
        table.rows[9].cells[1].text = '票务系统 (xs.bjstarfish.com)'
        
        self.doc.add_paragraph()
        
        # 适用范围
        self.doc.add_heading('1.2 适用范围', level=2)
        
        p = self.doc.add_paragraph()
        p.add_run('本手册适用于以下场景：')
        
        scenarios = [
            "租赁设备的初次安装部署",
            "客户端应用程序的版本更新",
            "设备更换后的应用重装",
            "系统故障后的应用恢复"
        ]
        
        for scenario in scenarios:
            p = self.doc.add_paragraph(f'• {scenario}', style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.5)
        
        self.doc.add_paragraph()
    
    def add_requirements(self):
        """添加安装要求"""
        self.doc.add_heading('二、安装要求', level=1)
        
        # 设备信息
        self.doc.add_heading('2.1 设备信息', level=2)
        
        p = self.doc.add_paragraph()
        p.add_run('设备品牌: ').font.bold = True
        p.add_run('Sunmi (商米)')
        
        p = self.doc.add_paragraph()
        p.add_run('设备型号: ').font.bold = True
        p.add_run('V2S')
        
        p = self.doc.add_paragraph()
        p.add_run('设备数量: ').font.bold = True
        p.add_run('15台')
        
        p = self.doc.add_paragraph()
        p.add_run('采购来源: ').font.bold = True
        p.add_run('亿思维科技有限公司')
        
        p = self.doc.add_paragraph()
        p.add_run('设备用途: ').font.bold = True
        p.add_run('租赁业务现场操作，包括租赁下单、设备归还、库存管理等')
        
        self.doc.add_paragraph()
        
        # 设备要求
        self.doc.add_heading('2.2 设备要求', level=2)
        
        requirements = [
            "Android系统版本：Android 5.0或以上",
            "存储空间：至少100MB可用空间",
            "网络连接：稳定的WiFi或4G/5G网络",
            "屏幕分辨率：建议1280x720或以上",
            "处理器：建议四核以上处理器"
        ]
        
        for req in requirements:
            p = self.doc.add_paragraph(f'• {req}')
            p.paragraph_format.left_indent = Inches(0.25)
        
        self.doc.add_paragraph()
        
        # 安装包要求
        self.doc.add_heading('2.3 安装包要求', level=2)
        
        p = self.doc.add_paragraph()
        p.add_run('安装包来源: ').font.bold = True
        p.add_run('由亿思维科技有限公司提供')
        
        p = self.doc.add_paragraph()
        p.add_run('文件格式: ').font.bold = True
        p.add_run('Android APK安装包')
        
        p = self.doc.add_paragraph()
        p.add_run('文件命名: ').font.bold = True
        p.add_run('租赁设备客户端_vX.X.X.apk（供应商提供最新版本）')
        
        self.doc.add_paragraph()
        
        # 重要提示
        p = self.doc.add_paragraph()
        p.add_run('重要提示: ').font.bold = True
        warning = p.add_run('请确保从供应商官方渠道获取安装包，不要使用来历不明的APK文件，以确保系统安全。')
        warning.font.color.rgb = RGBColor(255, 0, 0)
        warning.font.bold = True
        
        self.doc.add_paragraph()
    
    def add_installation_steps(self):
        """添加安装步骤"""
        self.doc.add_heading('三、安装步骤', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run('请按照以下步骤在租赁设备上安装客户端应用程序：')
        
        self.doc.add_paragraph()
        
        steps = [
            {
                "title": "步骤一：准备安装包",
                "content": "从供应商获取最新版本的Android APK安装包，并通过USB数据线、优盘或网络下载的方式将安装包传输到租赁设备上。",
                "note": "确保安装包文件完整，未损坏。建议将安装包放在设备的Download文件夹中。"
            },
            {
                "title": "步骤二：允许安装未知来源应用",
                "content": "进入设备的【设置】→【安全】→【未知来源】，开启允许安装未知来源应用的选项。部分设备需要在安装时临时授权。",
                "note": "不同Android版本的设置路径可能略有差异，请根据实际设备调整。"
            },
            {
                "title": "步骤三：安装应用程序",
                "content": "找到APK安装包文件，点击文件进行安装。按照系统提示，点击【安装】按钮，等待安装完成。",
                "note": "安装过程需要几秒到几十秒不等，请耐心等待。"
            },
            {
                "title": "步骤四：打开应用",
                "content": "安装完成后，点击【打开】按钮启动应用，或在桌面找到应用图标点击启动。",
                "note": "首次启动可能需要初始化，加载时间稍长。"
            },
            {
                "title": "步骤五：系统配置",
                "content": "首次启动后，需要配置系统连接信息，包括服务器地址、设备编号等。按照屏幕提示完成配置。",
                "note": "配置信息请向系统管理员或供应商技术人员获取。"
            },
            {
                "title": "步骤六：登录系统",
                "content": "使用分配的账号密码登录系统。登录后系统会自动同步租赁产品信息、库存数据等。",
                "note": "首次登录需要联网，确保网络连接正常。"
            },
            {
                "title": "步骤七：功能测试",
                "content": "登录成功后，测试以下功能：租赁下单、还设备、库存查询、订单查询等核心功能。",
                "note": "确保所有功能正常后再投入使用。"
            }
        ]
        
        for idx, step in enumerate(steps, 1):
            # 步骤标题
            heading = self.doc.add_heading(step['title'], level=2)
            heading_run = heading.runs[0]
            heading_run.font.color.rgb = RGBColor(68, 114, 196)
            
            # 操作内容
            p = self.doc.add_paragraph()
            p.add_run('操作说明: ').font.bold = True
            p.add_run(step['content'])
            
            # 注意事项
            if step.get('note'):
                p = self.doc.add_paragraph()
                p.add_run('注意事项: ').font.bold = True
                note_run = p.add_run(step['note'])
                note_run.font.color.rgb = RGBColor(255, 128, 0)
                note_run.font.italic = True
            
            self.doc.add_paragraph()
    
    def add_function_introduction(self):
        """添加功能介绍"""
        self.doc.add_heading('四、功能介绍', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run('租赁设备客户端主要功能包括：')
        
        self.doc.add_paragraph()
        
        functions = [
            {
                "name": "租赁下单",
                "desc": "支持现场租赁下单，选择租赁产品、数量、时长，生成租赁订单"
            },
            {
                "name": "还设备",
                "desc": "处理设备归还，记录归还时间，计算租赁费用，更新库存"
            },
            {
                "name": "库存管理",
                "desc": "实时查看设备库存情况，包括可租数量、已租数量、损坏数量等"
            },
            {
                "name": "订单查询",
                "desc": "查询租赁订单信息，包括订单状态、租赁时长、费用等"
            },
            {
                "name": "设备状态",
                "desc": "查看和更新设备状态，包括可用、已租、维修中等状态"
            },
            {
                "name": "数据同步",
                "desc": "与票务系统实时同步数据，确保信息一致性"
            }
        ]
        
        for func in functions:
            self.doc.add_heading(f'• {func["name"]}', level=2)
            p = self.doc.add_paragraph(func["desc"])
            p.paragraph_format.left_indent = Inches(0.25)
        
        self.doc.add_paragraph()
    
    def add_testing_checklist(self):
        """添加测试检查清单"""
        self.doc.add_heading('五、安装验收清单', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run('安装完成后，请逐项检查以下内容，确保所有功能正常：')
        
        self.doc.add_paragraph()
        
        checklist = [
            "□ 应用程序已成功安装",
            "□ 应用可以正常启动，无崩溃现象",
            "□ 网络连接正常，能访问服务器",
            "□ 登录功能正常，可以成功登录",
            "□ 租赁产品信息已同步",
            "□ 库存数据显示正确",
            "□ 租赁下单功能正常",
            "□ 还设备功能正常",
            "□ 订单查询功能正常",
            "□ 打印功能正常（如有）",
            "□ 扫码功能正常（如有）",
            "□ 应用运行流畅，无明显卡顿"
        ]
        
        for item in checklist:
            p = self.doc.add_paragraph(item)
            p.paragraph_format.left_indent = Inches(0.5)
            p.runs[0].font.size = Pt(11)
        
        self.doc.add_paragraph()
        
        # 验收签字
        p = self.doc.add_paragraph()
        p.add_run('验收人员: ').font.bold = True
        p.add_run('_______________    ')
        p.add_run('日期: ').font.bold = True
        p.add_run('_______________')
        
        self.doc.add_paragraph()
    
    def add_update_guide(self):
        """添加更新指南"""
        self.doc.add_heading('六、版本更新指南', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run('当供应商发布新版本时，可按以下步骤进行更新：')
        
        self.doc.add_paragraph()
        
        update_steps = [
            {
                "step": "获取新版本",
                "desc": "从供应商处获取最新版本的APK安装包"
            },
            {
                "step": "备份数据",
                "desc": "在更新前，确保重要数据已备份或已同步到服务器"
            },
            {
                "step": "卸载旧版本（可选）",
                "desc": "如果需要全新安装，可先卸载旧版本；如果仅升级，可直接安装覆盖"
            },
            {
                "step": "安装新版本",
                "desc": "按照第三章的安装步骤，安装新版本APK"
            },
            {
                "step": "重新登录",
                "desc": "安装完成后，重新登录系统"
            },
            {
                "step": "功能验证",
                "desc": "测试主要功能是否正常，确认数据同步完整"
            }
        ]
        
        for idx, step in enumerate(update_steps, 1):
            p = self.doc.add_paragraph()
            p.add_run(f'{idx}. {step["step"]}: ').font.bold = True
            p.add_run(step["desc"])
        
        self.doc.add_paragraph()
    
    def add_precautions(self):
        """添加注意事项"""
        self.doc.add_heading('七、注意事项', level=1)
        
        # 安装注意事项
        self.doc.add_heading('7.1 安装注意事项', level=2)
        
        install_notes = [
            "仅从供应商官方渠道获取安装包，确保应用安全",
            "安装前确认设备有足够的存储空间",
            "安装过程中保持设备电量充足，避免安装中断",
            "首次安装后需要进行完整的功能测试",
            "记录安装时间、版本号和操作人员"
        ]
        
        for note in install_notes:
            p = self.doc.add_paragraph(f'⚠ {note}')
            p.paragraph_format.left_indent = Inches(0.25)
            p.runs[0].font.color.rgb = RGBColor(255, 128, 0)
        
        self.doc.add_paragraph()
        
        # 使用注意事项
        self.doc.add_heading('7.2 使用注意事项', level=2)
        
        usage_notes = [
            "定期检查应用版本，及时更新到最新版本",
            "保持设备网络连接稳定，确保数据实时同步",
            "定期清理缓存，保持应用运行流畅",
            "妥善保管登录账号密码，不要泄露给无关人员",
            "发现异常及时联系技术支持，不要自行修改系统配置"
        ]
        
        for note in usage_notes:
            p = self.doc.add_paragraph(f'• {note}')
            p.paragraph_format.left_indent = Inches(0.25)
        
        self.doc.add_paragraph()
    
    def add_troubleshooting(self):
        """添加故障排查"""
        self.doc.add_heading('八、常见问题处理', level=1)
        
        problems = [
            {
                "problem": "安装失败",
                "solutions": [
                    "检查设备存储空间是否充足",
                    "确认已开启允许安装未知来源应用",
                    "检查APK文件是否完整未损坏",
                    "重启设备后重新安装"
                ]
            },
            {
                "problem": "无法启动应用",
                "solutions": [
                    "检查设备Android版本是否符合要求",
                    "清除应用缓存后重试",
                    "卸载后重新安装",
                    "联系供应商技术支持"
                ]
            },
            {
                "problem": "无法登录",
                "solutions": [
                    "检查网络连接是否正常",
                    "确认服务器地址配置正确",
                    "验证账号密码是否正确",
                    "联系管理员检查账号状态"
                ]
            },
            {
                "problem": "数据不同步",
                "solutions": [
                    "检查网络连接是否稳定",
                    "退出应用重新登录",
                    "手动触发数据同步",
                    "检查服务器是否正常运行"
                ]
            },
            {
                "problem": "应用运行卡顿",
                "solutions": [
                    "清理应用缓存和数据",
                    "关闭后台不必要的应用",
                    "重启设备",
                    "检查设备配置是否满足要求"
                ]
            }
        ]
        
        for idx, item in enumerate(problems, 1):
            self.doc.add_heading(f'8.{idx} {item["problem"]}', level=2)
            
            p = self.doc.add_paragraph()
            p.add_run('解决方案：').font.bold = True
            
            for solution in item["solutions"]:
                p = self.doc.add_paragraph(f'• {solution}')
                p.paragraph_format.left_indent = Inches(0.5)
            
            self.doc.add_paragraph()
    
    def add_maintenance(self):
        """添加维护说明"""
        self.doc.add_heading('九、日常维护', level=1)
        
        self.doc.add_heading('9.1 定期维护任务', level=2)
        
        maintenance_tasks = [
            {
                "frequency": "每日",
                "tasks": ["检查应用运行状态", "确认数据同步正常"]
            },
            {
                "frequency": "每周",
                "tasks": ["清理应用缓存", "检查设备存储空间", "测试主要功能"]
            },
            {
                "frequency": "每月",
                "tasks": ["检查应用版本", "备份重要数据", "进行全面功能测试"]
            }
        ]
        
        for task in maintenance_tasks:
            p = self.doc.add_paragraph()
            p.add_run(f'{task["frequency"]}维护: ').font.bold = True
            
            for item in task["tasks"]:
                p = self.doc.add_paragraph(f'• {item}', style='List Bullet')
                p.paragraph_format.left_indent = Inches(0.5)
        
        self.doc.add_paragraph()
        
        # 维护建议
        self.doc.add_heading('9.2 维护建议', level=2)
        
        suggestions = [
            "建立设备维护记录表，记录每次维护的时间和内容",
            "定期与供应商沟通，了解新版本发布信息",
            "培训操作人员，确保正确使用应用程序",
            "发现问题及时报告，不要等到影响业务时才处理",
            "保持与票务系统的对接正常，定期测试数据同步"
        ]
        
        for suggestion in suggestions:
            p = self.doc.add_paragraph(f'✓ {suggestion}')
            p.paragraph_format.left_indent = Inches(0.25)
        
        self.doc.add_paragraph()
    
    def add_contact(self):
        """添加技术支持"""
        self.doc.add_heading('十、技术支持', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run('如遇到技术问题或需要帮助，请联系供应商技术支持团队。')
        
        self.doc.add_paragraph()
        
        # 支持信息
        table = self.doc.add_table(rows=4, cols=2)
        table.style = 'Light Grid Accent 1'
        
        table.rows[0].cells[0].text = '项目'
        table.rows[0].cells[1].text = '信息'
        
        for cell in table.rows[0].cells:
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        table.rows[1].cells[0].text = '供应商'
        table.rows[1].cells[1].text = '亿思维科技有限公司'
        
        table.rows[2].cells[0].text = '支持内容'
        table.rows[2].cells[1].text = '应用安装、配置指导、故障排查、版本更新'
        
        table.rows[3].cells[0].text = '支持方式'
        table.rows[3].cells[1].text = '电话支持、远程协助、现场服务'
        
        self.doc.add_paragraph()
        
        # 文档信息
        p = self.doc.add_paragraph()
        p.add_run('文档版本: ').font.bold = True
        p.add_run('v1.0')
        
        p = self.doc.add_paragraph()
        p.add_run('发布日期: ').font.bold = True
        p.add_run('2023-11-05')
        
        p = self.doc.add_paragraph()
        p.add_run('适用版本: ').font.bold = True
        p.add_run('所有租赁设备客户端版本')
        
        self.doc.add_paragraph()
    
    def generate(self):
        """生成文档"""
        print("\n" + "="*60)
        print("生成租赁设备客户端发布手册")
        print("="*60)
        
        print("\n添加标题...")
        self.add_title()
        
        print("添加文档说明...")
        self.add_overview()
        
        print("添加安装要求...")
        self.add_requirements()
        
        print("添加安装步骤...")
        self.add_installation_steps()
        
        print("添加功能介绍...")
        self.add_function_introduction()
        
        print("添加安装验收清单...")
        self.add_testing_checklist()
        
        print("添加更新指南...")
        self.add_update_guide()
        
        print("添加注意事项...")
        self.add_precautions()
        
        print("添加故障排查...")
        self.add_troubleshooting()
        
        print("添加日常维护...")
        self.add_maintenance()
        
        print("添加技术支持...")
        self.add_contact()
        
        # 保存文档
        self.doc.save(OUTPUT_FILE)
        print(f"\n✓ 文档已生成: {OUTPUT_FILE}")
        print(f"  发布日期: 2023-11-05")
        print(f"  供应商: 亿思维科技有限公司")
        print(f"  设备品牌: Sunmi (商米)")
        print(f"  设备型号: V2S")
        print(f"  设备数量: 15台")
        print(f"  平台: Android")
        print(f"  内容: 安装步骤、功能介绍、维护指南")
        print("="*60 + "\n")
        
        return OUTPUT_FILE

def main():
    generator = RentalClientManualGenerator()
    generator.generate()

if __name__ == "__main__":
    main()
